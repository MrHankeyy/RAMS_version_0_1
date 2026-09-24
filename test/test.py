
"""根据 Excel 样本点计算目标响应和约束响应。

默认读取并原地更新脚本所在目录下的 ``ADAD_with_responses.xlsx``。
每次运行都会根据文件中当前存在的 X1、X2 样本重新计算响应。
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from openpyxl import load_workbook


def f_obj(x1: float, x2: float) -> float:
    """鲁棒优化校准目标：近边界窄深谷 + 远边界宽浅谷，目标为最小化。"""
    background = 20.0 + 0.03 * ((x1 - 2.0) ** 2 + (x2 - 2.0) ** 2)
    narrow = -8.0 * math.exp(
        -0.5 * ((x1 - 3.9) ** 2 + (x2 - 3.9) ** 2) / 0.18 ** 2
    )
    broad = -11.5 * math.exp(
        -0.5 * ((x1 - 2.0) ** 2 + (x2 - 2.0) ** 2) / 2.0 ** 2
    )
    return background + narrow + broad


def g_constr(x1: float, x2: float) -> float:
    """约束边界，满足 ``g(x) <= 0`` 的样本为可行样本。"""
    return x1 + x2 - 8.0


def _number(value, column_name: str, row_number: int) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"第 {row_number} 行的 {column_name} 不是有效数字：{value!r}"
        ) from exc


def _find_column(headers, candidates):
    normalized = {str(value).strip().lower(): value for value in headers}
    for candidate in candidates:
        value = normalized.get(candidate.lower())
        if value is not None:
            return value
    return None


def calculate_workbook(
    input_path: str | Path,
    output_path: str | Path | None = None,
    sheet_name: str | None = None,
) -> Path:
    """读取样本点、计算响应并写回 Excel。

    输入列支持 ``X1/x1`` 和 ``X2/x2``。已有 ``Y`` 列会被写入目标响应数值；
    如果没有该列，则创建 ``目标响应`` 列，同时创建 ``约束响应`` 列。
    """
    source = Path(input_path)
    if not source.exists():
        raise FileNotFoundError(f"找不到输入 Excel：{source}")
    destination = Path(output_path) if output_path else source
    workbook = load_workbook(source)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    if worksheet.max_row < 2:
        raise ValueError(f"工作表 {worksheet.title!r} 没有可计算的数据行")

    headers = [cell.value for cell in worksheet[1]]
    x1_column = _find_column(headers, ("X1",))
    x2_column = _find_column(headers, ("X2",))
    if x1_column is None or x2_column is None:
        raise ValueError("Excel 必须包含 X1 和 X2 样本列")

    objective_column = _find_column(headers, ("Y", "目标响应", "目标响应值"))
    if objective_column is None:
        objective_column = "目标响应"
        worksheet.cell(row=1, column=worksheet.max_column + 1, value=objective_column)

    constraint_column = _find_column(
        headers, ("约束响应", "约束响应值", "Constraint", "constraint")
    )
    if constraint_column is None:
        constraint_column = "约束响应"
        worksheet.cell(row=1, column=worksheet.max_column + 1, value=constraint_column)

    header_to_index = {
        str(cell.value).strip().lower(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    x1_index = header_to_index[str(x1_column).strip().lower()]
    x2_index = header_to_index[str(x2_column).strip().lower()]
    objective_index = header_to_index[str(objective_column).strip().lower()]
    constraint_index = header_to_index[str(constraint_column).strip().lower()]

    calculated_rows = 0
    for row_number in range(2, worksheet.max_row + 1):
        x1_value = worksheet.cell(row_number, x1_index).value
        x2_value = worksheet.cell(row_number, x2_index).value
        if x1_value is None and x2_value is None:
            continue
        x1 = _number(x1_value, "X1", row_number)
        x2 = _number(x2_value, "X2", row_number)
        worksheet.cell(row_number, objective_index, value=round(f_obj(x1, x2), 8))
        worksheet.cell(row_number, constraint_index, value=round(g_constr(x1, x2), 8))
        calculated_rows += 1

    workbook.save(destination)
    print(
        f"已计算 {calculated_rows} 个样本点；目标响应写入 {objective_column!r}，"
        f"约束响应写入 {constraint_column!r}。"
    )
    print(f"结果文件：{destination.resolve()}")
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(description="计算 Excel 样本点的目标与约束响应")
    default_input = Path(__file__).with_name("ADAD_with_responses.xlsx")
    parser.add_argument("input", nargs="?", default=str(default_input), help="输入 Excel 文件")
    parser.add_argument("--output", help="输出 Excel 文件，默认原地更新输入文件")
    parser.add_argument("--sheet", help="工作表名称，默认使用第一个工作表")
    args = parser.parse_args()
    calculate_workbook(args.input, args.output, args.sheet)


if __name__ == "__main__":
    main()
